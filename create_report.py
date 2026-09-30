import os
import subprocess
import sys

def install_and_import():
    try:
        import docx
    except ImportError:
        print("python-docx not installed. Installing...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
        import docx
    return docx

docx = install_and_import()
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Tiêu đề
title = doc.add_heading('PROJECT REPORT: PHASE 1 & 2 - NEW YORK CITY AIRBNB DATA ANALYSIS', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('THÔNG TIN TỔNG QUAN (OVERVIEW)', level=1)
doc.add_paragraph('- Tên dự án: Phân tích Dữ liệu Hệ thống Đặt phòng Airbnb tại Thành phố New York.')
doc.add_paragraph('- Dữ liệu sử dụng: AB_NYC_2019.csv (New York City Airbnb Open Data năm 2019).')

doc.add_heading('PHASE 1: DISCOVERY (TÌM HIỂU VÀ XÁC ĐỊNH BÀI TOÁN)', level=1)

doc.add_heading('1. Project Goals (Mục tiêu Dự án)', level=2)
doc.add_paragraph('Mục tiêu chính của dự án là phân tích tập dữ liệu Airbnb của thành phố New York để:')
doc.add_paragraph('• Hiểu rõ bức tranh toàn cảnh về thị trường cho thuê nhà ngắn hạn tại New York (sự phân bố, mức giá, mức độ phổ biến).', style='List Bullet')
doc.add_paragraph('• Khám phá các yếu tố chính ảnh hưởng đến giá thuê (price) và mức độ tương tác của khách hàng (number_of_reviews).', style='List Bullet')
doc.add_paragraph('• Đưa ra những insight (thông tin chi tiết) và khuyến nghị hữu ích cho các bên liên quan, giúp họ đưa ra quyết định kinh doanh hoặc lựa chọn chỗ ở tốt hơn.', style='List Bullet')

doc.add_heading('2. Stakeholders (Các bên liên quan)', level=2)
doc.add_paragraph('• Chủ nhà (Hosts): Những người muốn biết cách định giá phòng tối ưu, cách cải thiện dịch vụ để thu hút nhiều đánh giá tích cực và tăng tỷ lệ lấp đầy.', style='List Bullet')
doc.add_paragraph('• Khách du lịch (Guests / Travelers): Những người tìm kiếm thông tin để có thể đặt phòng ở khu vực an toàn, thuận tiện với mức giá hợp lý nhất.', style='List Bullet')
doc.add_paragraph('• Quản lý Nền tảng (Airbnb Management): Cần phân tích thị trường để đưa ra các chiến lược marketing phù hợp theo từng khu vực tại New York.', style='List Bullet')
doc.add_paragraph('• Chính quyền địa phương (City Planners / Regulators): Theo dõi ảnh hưởng của Airbnb lên tình hình bất động sản địa phương.', style='List Bullet')

doc.add_heading('3. Requirements (Yêu cầu của dự án)', level=2)
doc.add_paragraph('• Về mặt dữ liệu: Dữ liệu phải được làm sạch, xử lý các giá trị bị thiếu (missing values) và định dạng lại đúng kiểu (data types).', style='List Bullet')
doc.add_paragraph('• Về mặt phân tích: Thực hiện Phân tích Khám phá Dữ liệu (EDA) để tìm ra các xu hướng, mức giá trung bình theo khu vực, và sự khác biệt giữa các loại phòng (room types).', style='List Bullet')
doc.add_paragraph('• Về mặt trực quan hóa: Sử dụng các biểu đồ (Bar chart, Scatter plot, Heatmap) để hiển thị dữ liệu một cách dễ hiểu nhất cho các Stakeholders.', style='List Bullet')

doc.add_heading('4. Hypothesis (Các Giả thuyết cần kiểm chứng)', level=2)
doc.add_paragraph('• Giả thuyết 1: Manhattan là khu vực (neighbourhood_group) có giá thuê trung bình cao nhất so với các khu vực khác tại NYC.', style='List Bullet')
doc.add_paragraph('• Giả thuyết 2: Loại phòng "Entire home/apt" (Nguyên căn) mang lại lợi nhuận cao nhất nhưng lại có số đêm thuê tối thiểu (minimum_nights) yêu cầu dài hơn so với "Private room".', style='List Bullet')
doc.add_paragraph('• Giả thuyết 3: Các danh sách phòng có số lượng đánh giá (number_of_reviews) cao thường có mức giá vừa phải và tỷ lệ trống phòng (availability_365) thấp do được đặt liên tục.', style='List Bullet')

doc.add_heading('PHASE 2: DATA PREPARATION (CHUẨN BỊ DỮ LIỆU)', level=1)

doc.add_heading('1. Data Collection (Thu thập dữ liệu)', level=2)
doc.add_paragraph('• Nguồn dữ liệu: Dữ liệu được lấy từ nền tảng Kaggle (New York City Airbnb Open Data), phản ánh các hoạt động và số liệu công khai (public metrics) của Airbnb tại NYC trong năm 2019.', style='List Bullet')
doc.add_paragraph('• Hình thức: File định dạng .csv (AB_NYC_2019.csv) chứa 48,895 bản ghi (rows) và 16 trường dữ liệu (columns).', style='List Bullet')

doc.add_heading('2. Data Description (Mô tả Từ điển Dữ liệu)', level=2)
doc.add_paragraph('Dựa vào tập file được cung cấp, tập dữ liệu có các trường thông tin quan trọng sau:')
doc.add_paragraph('• id & name: Mã định danh và Tên của danh sách cho thuê.', style='List Bullet')
doc.add_paragraph('• host_id & host_name: Mã định danh và Tên của chủ nhà.', style='List Bullet')
doc.add_paragraph('• neighbourhood_group: Nhóm quận/khu vực lớn (VD: Manhattan, Brooklyn, Queens...).', style='List Bullet')
doc.add_paragraph('• neighbourhood: Khu vực lân cận cụ thể.', style='List Bullet')
doc.add_paragraph('• latitude & longitude: Tọa độ địa lý của phòng cho thuê.', style='List Bullet')
doc.add_paragraph('• room_type: Loại phòng cho thuê (Entire home/apt, Private room, Shared room).', style='List Bullet')
doc.add_paragraph('• price: Giá thuê một đêm ($).', style='List Bullet')
doc.add_paragraph('• minimum_nights: Số đêm thuê tối thiểu.', style='List Bullet')
doc.add_paragraph('• number_of_reviews & last_review: Tổng số lượt đánh giá và ngày đánh giá gần nhất.', style='List Bullet')
doc.add_paragraph('• reviews_per_month: Số lượt đánh giá trung bình mỗi tháng.', style='List Bullet')
doc.add_paragraph('• calculated_host_listings_count: Tổng số phòng mà chủ nhà đó đang cho thuê.', style='List Bullet')
doc.add_paragraph('• availability_365: Số ngày phòng trống/có sẵn trong năm (tối đa 365).', style='List Bullet')

doc.add_heading('3. Data Cleaning Strategy (Chiến lược làm sạch dữ liệu)', level=2)
doc.add_paragraph('Đây là các bước sẽ thực hiện (và báo cáo) trong công cụ phân tích (như Python, Excel, PowerBI):')
doc.add_paragraph('• Xử lý dữ liệu bị thiếu (Missing Values):', style='List Bullet')
doc.add_paragraph('   - Các cột last_review và reviews_per_month có thể có nhiều giá trị Null đối với những phòng chưa từng được đánh giá (number_of_reviews = 0). Cần thay thế (impute) bằng 0 hoặc loại bỏ tùy ngữ cảnh.')
doc.add_paragraph('   - Cột name hoặc host_name có thể thiếu 1 số ít dữ liệu. Có thể điền là "Unknown".')
doc.add_paragraph('• Xử lý dữ liệu bất thường (Outliers):', style='List Bullet')
doc.add_paragraph('   - Kiểm tra cột price (Giá phòng): Lọc bỏ các phòng có giá bằng 0 hoặc có giá quá cao (outliers phi lý).')
doc.add_paragraph('   - Kiểm tra cột minimum_nights: Lọc bỏ các yêu cầu thuê tối thiểu vượt quá 365 ngày vì đây là dữ liệu có thể bị lỗi.')
doc.add_paragraph('• Định dạng dữ liệu (Data Type Formatting): Đảm bảo cột last_review được định dạng dưới dạng chuỗi thời gian (DateTime).', style='List Bullet')

output_path = r'd:\dta\Project_Report_Phase_1_2.docx'
doc.save(output_path)
print(f"File DOCX created successfully at: {output_path}")
