import os
import subprocess
import sys

def install_and_import():
    try:
        import docx
    except ImportError:
        print("python-docx not installed. Installing...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
        import docx
    return docx

docx = install_and_import()
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Tiêu đề
title = doc.add_heading('PROJECT REPORT: NEW YORK CITY AIRBNB DATA ANALYSIS', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('THÔNG TIN TỔNG QUAN (OVERVIEW)', level=1)
doc.add_paragraph('- Tên dự án: Giải pháp Dữ liệu & Tối ưu Chiến lược Kinh doanh Airbnb tại New York.')
doc.add_paragraph('- Nguồn dữ liệu: Dữ liệu công khai từ Inside Airbnb (http://insideairbnb.com/get-the-data/) - Cụ thể: Thành phố New York (New York City), snapshot tháng 06/2026.')

doc.add_heading('PHASE 1: DISCOVERY (TÌM HIỂU VÀ XÁC ĐỊNH BÀI TOÁN)', level=1)

doc.add_heading('1. Business Objectives (Mục tiêu Cốt lõi & Bài toán Kinh doanh)', level=2)
doc.add_paragraph('Mục tiêu bao trùm của dự án là hoạt động như một "nhà tư vấn dữ liệu", cung cấp giải pháp giúp các Chủ nhà (Hosts) trên nền tảng Airbnb tại New York giải quyết 3 bài toán sống còn sau:')
doc.add_paragraph('• Bài toán Định giá (Pricing Strategy): Xây dựng mô hình Hồi quy (Regression) dự đoán giá thuê hợp lý nhất dựa trên vị trí, tiện ích và loại phòng, giúp chủ nhà tránh tình trạng đặt giá quá cao (ế khách) hoặc quá thấp (lỗ vốn).', style='List Bullet')
doc.add_paragraph('• Bài toán Định vị (Market Positioning): Dùng thuật toán Phân cụm (K-means Clustering) để nhóm các phòng thành các phân khúc đặc trưng, giúp chủ nhà xác định rõ đối thủ cạnh tranh và đối tượng khách hàng mục tiêu.', style='List Bullet')
doc.add_paragraph('• Bài toán Nâng cao Dịch vụ (Customer Satisfaction): Áp dụng NLP (Text/Sentiment Analysis) lên lịch sử bình luận để chỉ ra những yếu tố khiến khách hàng phật ý hoặc hài lòng nhất, từ đó cải thiện chất lượng dịch vụ để đạt đánh giá 5 sao.', style='List Bullet')

doc.add_heading('2. Stakeholders (Các bên liên quan)', level=2)
doc.add_paragraph('• Chủ nhà (Hosts) / Nhà đầu tư: Hưởng lợi trực tiếp từ các mô hình định giá và phân tích phản hồi để tối đa hóa lợi nhuận.', style='List Bullet')
doc.add_paragraph('• Khách du lịch (Guests): Dễ dàng tìm được chỗ ở phù hợp thông qua kết quả phân cụm thị trường.', style='List Bullet')

doc.add_heading('3. Requirements (Yêu cầu kỹ thuật)', level=2)
doc.add_paragraph('• Áp dụng chặt chẽ 6 giai đoạn của Data Analytics Lifecycle.', style='List Bullet')
doc.add_paragraph('• Làm sạch và tiền xử lý dữ liệu bảng (listings) và dữ liệu văn bản (reviews).', style='List Bullet')
doc.add_paragraph('• Ứng dụng thành thạo Machine Learning: Regression, Clustering và Text Analysis (TF-IDF).', style='List Bullet')

doc.add_heading('4. Hypothesis (Các giả thuyết nghiên cứu)', level=2)
doc.add_paragraph('• H1 (Về giá): Các phòng ở khu vực Manhattan và có thuộc tính "Entire home/apt" sẽ có mức giá trung bình cao nhất.')
doc.add_paragraph('• H2 (Về khách hàng): Phân khúc giá rẻ (dưới $100/đêm) sẽ có số lượng review (lượt khách) cao hơn đáng kể so với phân khúc cao cấp.')
doc.add_paragraph('• H3 (Về đánh giá): Các bình luận chứa nhiều từ khóa như "dirty", "loud", "small" sẽ tương quan mạnh mẽ với điểm đánh giá (review_scores_rating) thấp.')

doc.add_heading('PHASE 2: DATA PREPARATION (CHUẨN BỊ DỮ LIỆU)', level=1)

doc.add_heading('1. Data Collection (Thu thập dữ liệu)', level=2)
doc.add_paragraph('• Nguồn truy xuất: Toàn bộ dữ liệu được tải trực tiếp từ cổng thông tin mở http://insideairbnb.com/get-the-data/ (Chỉ trích xuất phần dữ liệu của riêng New York City, United States).', style='List Bullet')
doc.add_paragraph('• Các tập tin chính được giữ lại phục vụ dự án: listings.csv (đặc tả phòng) và reviews.csv (bình luận).', style='List Bullet')

doc.add_heading('2. Data Description (Mô tả Dữ liệu)', level=2)
doc.add_paragraph('• Tập listings.csv: Chứa các thuộc tính định lượng và phân loại (id, host_name, neighbourhood, room_type, price, minimum_nights, reviews_per_month...).', style='List Bullet')
doc.add_paragraph('• Tập reviews.csv: Chứa dữ liệu thời gian và văn bản thô (date, comments).', style='List Bullet')

doc.add_heading('3. Data Cleaning Strategy (Chiến lược làm sạch dữ liệu)', level=2)
doc.add_paragraph('• Xử lý Dị biệt (Outliers): Khảo sát và loại bỏ các phòng có mức giá (price) bất hợp lý ($0 hoặc quá đắt) và số đêm tối thiểu (minimum_nights) vượt ngoài thực tế.', style='List Bullet')
doc.add_paragraph('• Xử lý Giá trị Khuyết (Missing Values): Xử lý các dòng thiếu thông tin đánh giá bằng cách thay thế (imputation) hoặc loại bỏ (drop).', style='List Bullet')
doc.add_paragraph('• Text Processing: Làm sạch cột comments (loại bỏ stop words, dấu câu) để chuẩn bị ma trận TF-IDF cho NLP.', style='List Bullet')

output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'report', 'Project_Report_Phase_1_2.docx')
doc.save(output_path)
print(f"File DOCX created successfully at: {output_path}")
