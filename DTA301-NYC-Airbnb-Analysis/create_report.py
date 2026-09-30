import os
import subprocess
import sys

def install_and_import():
    try:
        import docx
    except ImportError:
        print("python-docx not installed. Installing...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
        import docx
    return docx

docx = install_and_import()
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Tiêu đề
title = doc.add_heading('PROJECT REPORT: NEW YORK CITY AIRBNB DATA ANALYSIS', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('THÔNG TIN TỔNG QUAN (OVERVIEW)', level=1)
doc.add_paragraph('- Tên dự án: Phân tích Dữ liệu Hệ thống Đặt phòng Airbnb tại Thành phố New York.')
doc.add_paragraph('- Dữ liệu sử dụng: Tập dữ liệu công khai Inside Airbnb (Năm 2024) bao gồm listings.csv và reviews.csv.')

doc.add_heading('PHASE 1: DISCOVERY (TÌM HIỂU VÀ XÁC ĐỊNH BÀI TOÁN)', level=1)

doc.add_heading('1. Project Goals (Mục tiêu Dự án)', level=2)
doc.add_paragraph('Mục tiêu chính của dự án là phân tích tập dữ liệu Airbnb tại New York kết hợp với vòng đời phân tích dữ liệu (Data Analytics Lifecycle) để:')
doc.add_paragraph('• Hiểu rõ bức tranh thị trường cho thuê ngắn hạn tại New York sau đại dịch và tác động của các chính sách mới (Local Law 18).', style='List Bullet')
doc.add_paragraph('• Dự đoán giá thuê (Price Prediction) thông qua mô hình Hồi quy (Regression).', style='List Bullet')
doc.add_paragraph('• Phân cụm khách hàng/phòng (Clustering) để tìm ra các phân khúc nhà Airbnb đặc trưng.', style='List Bullet')
doc.add_paragraph('• Phân tích cảm xúc (Sentiment Analysis / Text Analysis) trên dữ liệu bình luận của khách hàng.', style='List Bullet')

doc.add_heading('2. Stakeholders (Các bên liên quan)', level=2)
doc.add_paragraph('• Chủ nhà (Hosts): Biết cách định giá phòng tối ưu và cải thiện dịch vụ dựa trên phản hồi.', style='List Bullet')
doc.add_paragraph('• Khách du lịch (Guests): Tìm được chỗ ở phù hợp với nhu cầu và túi tiền.', style='List Bullet')

doc.add_heading('3. Requirements (Yêu cầu của dự án theo đề cương)', level=2)
doc.add_paragraph('• Tuân thủ 6 giai đoạn của Data Analytics Lifecycle.', style='List Bullet')
doc.add_paragraph('• Về dữ liệu: Làm sạch, xử lý dữ liệu thiếu và chuẩn bị dữ liệu văn bản (text data) từ reviews.', style='List Bullet')
doc.add_paragraph('• Về mô hình: Ứng dụng K-means (Clustering), Linear Regression và TF-IDF/Sentiment Analysis.', style='List Bullet')

doc.add_heading('PHASE 2: DATA PREPARATION (CHUẨN BỊ DỮ LIỆU)', level=1)

doc.add_heading('1. Data Collection (Thu thập dữ liệu)', level=2)
doc.add_paragraph('• Nguồn dữ liệu: Inside Airbnb (cập nhật mới nhất 2024).', style='List Bullet')
doc.add_paragraph('• Các tập tin chính: listings.csv (thông tin phòng/giá) và reviews.csv (lịch sử đánh giá).', style='List Bullet')

doc.add_heading('2. Data Description (Mô tả Dữ liệu)', level=2)
doc.add_paragraph('• Tập listings.csv: Chứa các thuộc tính như id, host_name, neighbourhood_group, room_type, price, minimum_nights, number_of_reviews...', style='List Bullet')
doc.add_paragraph('• Tập reviews.csv: Chứa listing_id, date (ngày đánh giá), và comments (nội dung đánh giá text).', style='List Bullet')

doc.add_heading('3. Data Cleaning Strategy (Chiến lược làm sạch dữ liệu)', level=2)
doc.add_paragraph('• Xử lý Outliers: Lọc bỏ các phòng có price quá vô lý (bằng 0 hoặc cao đột biến).', style='List Bullet')
doc.add_paragraph('• Xử lý Missing Values: Điền khuyết hoặc loại bỏ các dòng bị thiếu tọa độ, giá trị đánh giá.', style='List Bullet')
doc.add_paragraph('• Text Processing: Làm sạch cột comments (loại bỏ stop words, dấu câu) để chuẩn bị cho NLP.', style='List Bullet')
doc.add_paragraph('• Feature Engineering: Trích xuất các thuộc tính thời gian (năm/tháng) từ cột date trong reviews.csv để làm Time Series Analysis.', style='List Bullet')

output_path = r'd:\dta\Project_Report_Phase_1_2.docx'
doc.save(output_path)
print(f"File DOCX created successfully at: {output_path}")
